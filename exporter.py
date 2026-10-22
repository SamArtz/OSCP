"""XLSX estándar de compras, con cantidades confirmadas y auditoría de ajustes."""

from datetime import date
from io import BytesIO
from typing import Annotated

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, ConfigDict, Field, StringConstraints

from engine import ReplenishmentConfig, money


class ForecastOptions(BaseModel):
    model_config = ConfigDict(extra="forbid")
    ventana_dias: int = Field(default=60, ge=30, le=60, strict=True)
    nivel_servicio: float = Field(default=0.95, gt=0.5, lt=1)
    peso_estacional: float = Field(default=0.5, ge=0, le=1)
    dias_revision: int = Field(default=0, ge=0, le=30, strict=True)

    def engine_config(self):
        return ReplenishmentConfig(self.ventana_dias, self.nivel_servicio, self.peso_estacional, self.dias_revision)


Identifier = Annotated[str, StringConstraints(strip_whitespace=True, min_length=1, max_length=200)]


class ConfirmedLine(BaseModel):
    model_config = ConfigDict(extra="forbid")
    proveedor_id: Identifier
    sku_id: Identifier
    cantidad_cajas: int = Field(ge=0, le=1_000_000, strict=True)


class PurchaseExport(BaseModel):
    model_config = ConfigDict(extra="forbid")
    fecha_corte: date
    huella_datos: str = Field(pattern=r"^[a-f0-9]{64}$")
    parametros: ForecastOptions = Field(default_factory=ForecastOptions)
    lineas: list[ConfirmedLine] = Field(min_length=1, max_length=10_000)


def workbook_bytes(lines: list[tuple[dict, int]], order: PurchaseExport) -> bytes:
    """Los valores de texto se escriben como strings, nunca como fórmulas Excel."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Orden_Compra"
    sheet.append([
        "proveedor_id", "sku_id", "sku_proveedor", "descripcion", "cantidad_cajas",
        "unidades_por_caja", "cantidad_unidades", "costo_unitario_usd", "total_usd",
    ])
    audit = workbook.create_sheet("Auditoria")
    audit.append(["proveedor_id", "sku_id", "sugerencia_cajas", "confirmado_cajas", "diferencia_cajas"])
    for row, boxes in sorted(lines, key=lambda pair: (pair[0]["proveedor_id"], pair[0]["sku_id"])):
        units = boxes * row["unidades_por_caja"]
        sheet.append([
            row["proveedor_id"], row["sku_id"], row["sku_proveedor"], row["nombre_producto"], boxes,
            row["unidades_por_caja"], units, row["costo_unitario_usd"], money(units, row["costo_unitario_usd"]),
        ])
        audit.append([row["proveedor_id"], row["sku_id"], row["cantidad_sugerida_cajas"], boxes, boxes - row["cantidad_sugerida_cajas"]])
    metadata = workbook.create_sheet("Parametros")
    metadata.append(["parametro", "valor"])
    metadata.append(["formato", "XLSX estándar v1; no es una integración certificada con un ERP"])
    metadata.append(["fecha_corte", order.fecha_corte.isoformat()])
    metadata.append(["huella_datos", order.huella_datos])
    for key, value in order.parametros.model_dump().items():
        metadata.append([key, value])
    for tab in workbook:
        tab.freeze_panes = "A2"
        tab.auto_filter.ref = tab.dimensions
        for cell in tab[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="153B38")
        for cells in tab.iter_rows(min_row=2):
            for cell in cells:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top")
        for column in tab.columns:
            length = max(len(str(cell.value or "")) for cell in column)
            tab.column_dimensions[column[0].column_letter].width = min(60, max(18, length + 2))
    for cells in sheet.iter_rows(min_row=2, min_col=8, max_col=9):
        for cell in cells:
            cell.number_format = '"$"#,##0.00'
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
