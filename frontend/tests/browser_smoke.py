"""Prueba con navegador y API reales; requiere ambos servidores y los tres Excel.

Opcional: pip install playwright; playwright install chromium.
CHROMIUM_PATH permite usar un Chromium ya instalado. No genera compras reales.
"""

import os
import re
from pathlib import Path
import shutil
import tempfile
from decimal import Decimal, ROUND_HALF_UP

from openpyxl import load_workbook
from playwright.sync_api import expect, sync_playwright


def run():
    root = Path(__file__).resolve().parents[2]
    output = Path(os.getenv("BROWSER_OUTPUT_DIR", str(Path(tempfile.gettempdir()) / "ocs-browser-check")))
    output.mkdir(parents=True, exist_ok=True)
    chromium = os.getenv("CHROMIUM_PATH") or shutil.which("chromium")
    with sync_playwright() as playwright:
        options = {"headless": True}
        if chromium:
            options["executable_path"] = chromium
        browser = playwright.chromium.launch(**options)
        try:
            page = browser.new_page(viewport={"width": 1440, "height": 1000}, accept_downloads=True)
            errors = []
            page.on("pageerror", lambda error: errors.append(str(error)))
            with page.expect_response(lambda response: "/api/replenishment?" in response.url, timeout=60000) as received:
                page.goto(os.getenv("FRONTEND_URL", "http://127.0.0.1:3000") + "/replenishment")
            data = received.value.json()
            assert data["estado"] == "exito", data
            expect(page.get_by_text("Datos validados", exact=True)).to_be_visible()
            dashboard = page.request.get(os.getenv("FRONTEND_URL", "http://127.0.0.1:3000") + "/api/dashboard?ventana_dias=60").json()
            assert data["huella_datos"] == dashboard["huella_datos"]
            assert data["salud_catalogo"] == dashboard["salud_catalogo"]
            expected_risk = data["salud_catalogo"]["quiebre"]
            metric = page.locator(".metric-card").filter(has_text="RIESGO DE QUIEBRE")
            expect(metric.locator(".metric-value")).to_have_text(str(expected_risk))
            expect(page.locator(".attention-strip")).to_contain_text(f"{expected_risk} SKU en quiebre o riesgo a 14 días")
            search = page.get_by_label("Buscar SKU, descripción o marca")
            for sku in ["SKU-1009", "SKU-1029"]:
                source = next((row for row in data["orden_sugerida"] if row["sku_id"] == sku), None)
                if source is None:
                    continue
                search.fill(sku)
                expect(page.locator("tbody tr")).to_have_count(1)
                expect(page.locator("tbody td").nth(4).locator("span").first).to_have_text(f"{source['stock_disponible']:,}")
                expect(page.locator("tbody td").nth(5).locator("span").first).to_have_text(f"{source['stock_en_transito']:,}")
                expect(page.locator("tbody td").nth(6).locator("span").first).to_have_attribute("title", re.compile("sum\\(w_t"))
            search.fill("")
            page.get_by_role("button", name="Riesgo de quiebre", exact=False).click()
            expect(page.locator("tbody tr")).to_have_count(min(expected_risk, 10))
            assert all("Quiebre / riesgo" in text for text in page.locator("tbody .status-badge").all_text_contents())
            page.get_by_role("button", name="Todos", exact=False).click()
            print("PASS: salud OCS/dashboard coherente, banner/filtro/badges, SKU-1009/1029 y fórmula auditable", flush=True)
            first = next(field for field in page.locator('input[type="number"]:enabled').all() if int(field.input_value()) > 0)
            sku = first.get_attribute("aria-label").removeprefix("Cajas confirmadas ")
            source_row = next(row for row in data["orden_sugerida"] if row["sku_id"] == sku)
            original = int(first.input_value())
            if len(data["orden_sugerida"]) > 10:
                page.get_by_role("button", name="Página siguiente").click()
                expect(page.locator(".pagination")).to_contain_text("2 /")
                page.get_by_role("button", name="Página anterior").click()
            supplier = data["orden_sugerida"][0]["proveedor_id"]
            page.get_by_label("Filtrar por proveedor").select_option(supplier)
            assert all(supplier in text for text in page.locator("tbody .sku-cell").all_text_contents())
            page.get_by_label("Filtrar por proveedor").select_option("")
            page.get_by_label("Filtrar por categoría").select_option(source_row["categoria"])
            page.get_by_label("Filtrar por marca").select_option(source_row["marca"])
            assert all(source_row["marca"] in text for text in page.locator("tbody .description-cell").all_text_contents())
            page.get_by_role("button", name="Limpiar", exact=False).click()
            page.keyboard.press("/")
            search = page.get_by_label("Buscar SKU, descripción o marca")
            expect(search).to_be_focused()
            search.fill(sku)
            field = page.get_by_role("spinbutton", name=f"Cajas confirmadas {sku}", exact=True)
            field.fill("1" if source_row["moq_cajas"] > 1 else "-1")
            expect(field).to_have_attribute("aria-invalid", "true")
            export = page.get_by_role("button", name="Descargar Orden de Compra (.xlsx)", exact=True)
            expect(export).to_be_disabled()
            page.get_by_role("button", name=f"Restaurar sugerencia {sku}", exact=True).click()
            expect(field).to_have_value(str(original))
            field.fill(str(original + 1))
            expect(export).to_be_enabled()
            page.get_by_role("button", name=f"Ver alertas de {sku}", exact=True).click()
            expect(page.get_by_role("dialog")).to_be_visible()
            expect(page.get_by_text("Prioridad de rotación FEFO")).to_be_visible()
            page.keyboard.press("Escape")
            search.fill("__sin_coincidencias__")  # todas las líneas quedan ocultas
            expect(page.locator("tbody tr")).to_have_count(0)
            with page.expect_download(timeout=60000) as downloaded:
                export.click()
            target = output / "compra_confirmada.xlsx"
            downloaded.value.save_as(target)
            book = load_workbook(target)
            lines = list(book["Orden_Compra"].values)[1:]
            catalog_costs = {row["sku_id"]: Decimal(str(row["costo_unitario_usd"])) for row in data["orden_sugerida"]}
            exported_total = Decimal(0)
            for line in lines:
                assert Decimal(str(line[7])) == catalog_costs[line[1]]
                expected_cost = (Decimal(line[6]) * catalog_costs[line[1]]).quantize(Decimal(".01"), rounding=ROUND_HALF_UP)
                assert Decimal(str(line[8])) == expected_cost
                exported_total += expected_cost
            expect(page.locator(".summary-total strong")).to_have_text(f"${exported_total:,.2f}")
            assert next(line for line in lines if line[1] == sku)[4] == original + 1
            expected = sum(row["cantidad_sugerida_cajas"] > 0 and row["estado_sugerencia"] == "CALCULADO" for row in data["orden_sugerida"])
            assert len(lines) == expected
            print("PASS: filtros, búsqueda, paginación, MOQ, ajuste, FEFO y XLSX con selección global", flush=True)
            search.fill("")
            page.get_by_role("button", name="Activar modo oscuro").click()
            expect(page.locator("html")).to_have_class("dark")
            page.screenshot(path=str(output / "oscuro.png"), full_page=True)
            page.get_by_role("button", name="Activar modo claro").click()
            page.get_by_role("button", name="Cargar archivos", exact=True).click()
            for name in ["ventas_historicas", "catalogo_sku", "inventario_actual"]:
                page.locator("#" + name).set_input_files(str(root / f"{name}.xlsx"))
            page.once("dialog", lambda dialog: dialog.accept())
            page.get_by_role("button", name="Calcular sugerencias", exact=True).click()
            expect(page.get_by_role("dialog")).to_have_count(0)
            expect(page.get_by_text("Datos validados", exact=True)).to_be_visible(timeout=60000)
            expect(page.get_by_role("spinbutton", name=f"Cajas confirmadas {sku}", exact=True)).to_have_value(str(original))
            with page.expect_download(timeout=60000) as uploaded_download:
                export.click()
            uploaded_download.value.save_as(output / "compra_uploads.xlsx")
            assert len(list(load_workbook(output / "compra_uploads.xlsx")["Orden_Compra"].values)) - 1 == expected
            page.screenshot(path=str(output / "escritorio.png"), full_page=True)
            print("PASS: carga y exportación de los tres Excel desde el navegador", flush=True)
            page.set_viewport_size({"width": 390, "height": 844})
            page.screenshot(path=str(output / "movil.png"), full_page=True)
            assert page.evaluate("document.documentElement.scrollWidth <= innerWidth")
            page.get_by_role("button", name="Cargar archivos", exact=True).click()
            expect(page.get_by_role("dialog")).to_be_visible()
            page.keyboard.press("Escape")
            page.route("**/api/replenishment?*", lambda route: route.fulfill(status=502, content_type="application/json", body='{"errores":[{"mensaje":"Servicio de prueba no disponible"}]}'))
            page.get_by_role("button", name="Recalcular sugerencias", exact=True).click()
            expect(page.get_by_role("alert").filter(has_text="Servicio de prueba no disponible")).to_be_visible()
            expect(export).to_be_disabled()
            assert page.locator("tbody tr").count() > 0  # cálculo anterior visible
            assert not errors, errors
            print("PASS: modo claro/oscuro, móvil sin desbordamiento, diálogo y fallo de conexión", flush=True)
        finally:
            browser.close()


if __name__ == "__main__":
    run()
