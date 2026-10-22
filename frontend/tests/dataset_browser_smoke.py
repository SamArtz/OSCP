"""Verifica que uploads y ventana viajan juntos entre OCS y dashboard.

Usa copias temporales de los tres Excel; requiere ambos servidores y Playwright.
"""

import os
from pathlib import Path
import shutil
import tempfile

from openpyxl import load_workbook
from playwright.sync_api import expect, sync_playwright


def run():
    root = Path(__file__).resolve().parents[2]
    base = os.getenv("FRONTEND_URL", "http://127.0.0.1:3000")
    with tempfile.TemporaryDirectory() as temporary:
        directory = Path(temporary)
        for name in ["ventas_historicas", "catalogo_sku", "inventario_actual"]:
            shutil.copyfile(root / f"{name}.xlsx", directory / f"{name}.xlsx")
        # Una diferencia deliberada hace detectable cualquier retorno a los archivos locales.
        for name, sheet, field, value in [
            ("catalogo_sku", "Catalogo", "costo_unitario_usd", 3.21),
            ("inventario_actual", "Inventario", "stock_disponible", 100000),
        ]:
            path = directory / f"{name}.xlsx"
            workbook = load_workbook(path)
            table = workbook[sheet]
            headers = [cell.value for cell in table[1]]
            sku_column, value_column = headers.index("sku_id") + 1, headers.index(field) + 1
            for cells in table.iter_rows(min_row=2):
                if cells[sku_column - 1].value == "SKU-1000":
                    cells[value_column - 1].value = value
            workbook.save(path)
        with sync_playwright() as playwright:
            chromium = os.getenv("CHROMIUM_PATH") or shutil.which("chromium")
            browser = playwright.chromium.launch(headless=True, **({"executable_path": chromium} if chromium else {}))
            try:
                page = browser.new_page(viewport={"width": 1440, "height": 1000})
                errors = []
                page.on("pageerror", lambda error: errors.append(str(error)))
                with page.expect_response(lambda response: "/api/replenishment?" in response.url, timeout=60000) as first:
                    page.goto(base + "/replenishment")
                original = first.value.json()
                expect(page.get_by_text("Datos validados", exact=True)).to_be_visible()
                page.get_by_role("button", name="Cargar archivos", exact=True).click()
                for name in ["ventas_historicas", "catalogo_sku", "inventario_actual"]:
                    page.locator("#" + name).set_input_files(str(directory / f"{name}.xlsx"))
                with page.expect_response(lambda response: "/api/replenishment?" in response.url, timeout=60000) as uploaded:
                    page.get_by_role("button", name="Calcular sugerencias", exact=True).click()
                changed = uploaded.value.json()
                expect(page.get_by_text("Datos validados", exact=True)).to_be_visible()
                require_sku = next(row for row in changed["orden_sugerida"] if row["sku_id"] == "SKU-1000")
                assert require_sku["stock_disponible"] == 100000
                assert require_sku["costo_unitario_usd"] == 3.21
                assert changed["huella_datos"] != original["huella_datos"]
                with page.expect_response(lambda response: "/api/dashboard?" in response.url, timeout=60000) as received:
                    page.get_by_role("link", name="Tablero ejecutivo", exact=True).click()
                dashboard = received.value.json()
                assert received.value.request.method == "POST"
                assert dashboard["huella_datos"] == changed["huella_datos"]
                assert dashboard["salud_catalogo"] == changed["salud_catalogo"]
                expect(page.get_by_text("3 archivos cargados", exact=True)).to_be_visible()
                with page.expect_response(lambda response: "/api/dashboard?ventana_dias=30" in response.url, timeout=60000) as shorter:
                    page.get_by_label("Ventana del tablero").select_option("30")
                dashboard_30 = shorter.value.json()
                expect(page.get_by_text("Datos validados", exact=True)).to_be_visible()
                with page.expect_response(lambda response: "/api/replenishment?ventana_dias=30" in response.url, timeout=60000) as returned:
                    page.get_by_role("link", name="Mesa de compras", exact=True).click()
                order_30 = returned.value.json()
                assert order_30["huella_datos"] == changed["huella_datos"] == dashboard_30["huella_datos"]
                assert order_30["salud_catalogo"] == dashboard_30["salud_catalogo"]
                assert order_30["parametros"]["ventana_dias"] == 30
                expect(page.get_by_text("3 archivos cargados", exact=True)).to_be_visible()
                # Recargar termina la selección en memoria y vuelve a los originales locales.
                with page.expect_response(lambda response: "/api/replenishment?ventana_dias=60" in response.url, timeout=60000) as reloaded:
                    page.reload()
                assert reloaded.value.json()["huella_datos"] == original["huella_datos"]
                assert not errors, errors
                print("PASS: archivos distintos conservados entre pantallas, misma salud y ventana, recarga descarta uploads; originales intactos", flush=True)
            finally:
                browser.close()


if __name__ == "__main__":
    run()
