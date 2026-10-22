"""Conciliación reproducible de los tres Excel, API y XLSX, sin cambiar fuentes.

Ejecutar: python audit.py --data-dir . --output-dir reports
Solo necesita requirements.txt; no requiere servidores ni datos comerciales en Git.
"""

import argparse
from collections import Counter
from decimal import Decimal, ROUND_HALF_UP
import hashlib
from io import BytesIO
import json
from math import ceil, isclose, sqrt
from pathlib import Path
from statistics import NormalDist

from openpyxl import Workbook, load_workbook
import pandas as pd

import main
from engine import ReplenishmentConfig
from exporter import PurchaseExport


class AuditFailure(ValueError):
    """Una salida del sistema no coincide con el archivo o con su contrato."""


def require(condition, message):
    if not condition:
        raise AuditFailure(message)


def usd(value):
    return Decimal(str(value)).quantize(Decimal(".01"), rounding=ROUND_HALF_UP)


def source_summary():
    return {name: {"origen": "local", "filas_leidas": None} for name in main.SCHEMAS}


def response_json(response):
    body = json.loads(response.body)
    require(response.status_code == 200 and body.get("estado") == "exito", f"Validación rechazada: {body.get('errores')}")
    return body


def reconcile(contents: dict[str, bytes], config: ReplenishmentConfig | None = None) -> tuple[dict, list[dict], bytes]:
    """Lee los operandos directamente del Excel, aparte de los DataFrames del motor."""
    config = config or ReplenishmentConfig()
    original_hashes = {name: hashlib.sha256(content).hexdigest() for name, content in contents.items()}
    etl = response_json(main.validate_contents(contents, source_summary(), [], 5))
    order = response_json(main.replenishment_contents(contents, source_summary(), [], config))
    dashboard = response_json(main.dashboard_contents(contents, source_summary(), [], config))
    raw = {name: pd.read_excel(BytesIO(contents[name]), sheet_name=schema["sheet"], keep_default_na=False)
           for name, schema in main.SCHEMAS.items()}
    catalog = raw["catalogo_sku"].copy()
    catalog.columns = catalog.columns.str.strip().str.lower()
    catalog["sku_id"] = catalog.sku_id.astype(str).str.strip()
    require(not catalog.sku_id.duplicated().any(), "SKU duplicados en catálogo")
    catalog = catalog.set_index("sku_id", verify_integrity=True)
    inventory, sales = raw["inventario_actual"].copy(), raw["ventas_historicas"].copy()
    for frame in [inventory, sales]:
        frame.columns = frame.columns.str.strip().str.lower()
        frame["sku_id"] = frame.sku_id.astype(str).str.strip()
    cutoff = pd.Timestamp(order["fecha_corte"])
    inventory["expiry"] = pd.to_datetime(inventory.fecha_vencimiento_lote).dt.normalize()
    sales["day"] = pd.to_datetime(sales.fecha_factura).dt.normalize()
    invoices = sales.loc[sales.tipo_doc.str.strip().str.upper().eq("FACTURA") & sales.day.lt(cutoff)].copy()
    invoices["physical"] = invoices.unidades_vendidas + invoices.unidades_bonificadas
    expected_skus = set(catalog.index)
    rows = order["orden_sugerida"]
    detail = dashboard["detalle_skus"]
    require(len(rows) == len(expected_skus) == len({row["sku_id"] for row in rows}), "SKU repetidos o faltantes en OCS")
    require(set(row["sku_id"] for row in rows) == expected_skus, "OCS no representa exactamente el catálogo")
    require(len(detail) == len(expected_skus) == len({row["sku_id"] for row in detail}), "SKU repetidos o faltantes en dashboard")
    require(set(row["sku_id"] for row in detail) == expected_skus, "Dashboard no representa exactamente el catálogo")
    dashboard_by_sku = {row["sku_id"]: row for row in detail}
    require(order["huella_datos"] == dashboard["huella_datos"], "Las pantallas usaron versiones distintas de los archivos")
    verified, total = [], Decimal(0)
    for row in rows:
        sku = row["sku_id"]
        item, health = catalog.loc[sku], dashboard_by_sku[sku]
        lots = inventory.loc[inventory.sku_id.eq(sku)]
        available = int(lots.stock_disponible.sum()) if len(lots) else None
        transit = int(lots.stock_en_transito.sum()) if len(lots) else None
        usable = int(lots.loc[lots.expiry.ge(cutoff), "stock_disponible"].sum()) if len(lots) else None
        for name, expected in [("stock_disponible", available), ("stock_en_transito", transit), ("stock_utilizable", usable)]:
            require(row[name] == health[name] == expected, f"{sku}: {name} no coincide con inventario")
        require(row["nombre_producto"] == str(item.nombre_producto).strip(), f"{sku}: descripción cruzada")
        require(row["proveedor_id"] == str(item.proveedor_id).strip(), f"{sku}: proveedor cruzado")
        require(Decimal(str(row["costo_unitario_usd"])) == Decimal(str(item.costo_unitario_usd)), f"{sku}: costo unitario ajeno al catálogo")
        units, boxes, case_size = row["cantidad_sugerida_unidades"], row["cantidad_sugerida_cajas"], int(item.unidades_por_caja)
        require(units == boxes * case_size and boxes >= 0 and int(boxes) == boxes, f"{sku}: unidades/cajas inconsistentes")
        require(boxes == 0 or boxes >= int(item.moq_cajas), f"{sku}: MOQ incumplido")
        expected_cost = usd(Decimal(units) * Decimal(str(item.costo_unitario_usd)))
        require(usd(row["costo_total_usd"]) == expected_cost, f"{sku}: importe distinto de unidades × costo catálogo")
        total += expected_cost
        movement = invoices.loc[invoices.sku_id.eq(sku)]
        mean, forecast, safety, rop = None, None, None, None
        if len(movement):
            start = max(cutoff - pd.Timedelta(days=config.window_days), movement.day.min())
            dates = pd.date_range(start, cutoff - pd.Timedelta(days=1))
            daily = movement.groupby("day").physical.sum().to_dict()
            numerator = sum((offset + 1) * float(daily.get(day, 0)) for offset, day in enumerate(dates))
            denominator = len(dates) * (len(dates) + 1) // 2
            mean = numerator / denominator
            audit = row["auditoria_demanda"]
            require(audit == health["auditoria_demanda"], f"{sku}: auditoría de demanda distinta entre pantallas")
            require(audit["suma_pesos"] == denominator and isclose(audit["suma_demanda_por_peso"], numerator), f"{sku}: pesos de demanda no reproducibles")
            require(isclose(audit["media_diaria_sin_redondear"], mean, abs_tol=1e-12), f"{sku}: media no coincide con ventas")
            weekday_counts = Counter(day.dayofweek for day in dates)
            factors = [1.0] * 7
            if mean > 0 and config.seasonal_weight > 0 and all(weekday_counts[day] >= 4 for day in range(7)):
                weekday_means = [
                    sum((offset + 1) * float(daily.get(day, 0)) for offset, day in enumerate(dates) if day.dayofweek == weekday)
                    / sum(offset + 1 for offset, day in enumerate(dates) if day.dayofweek == weekday)
                    for weekday in range(7)
                ]
                level = sum(weekday_means) / 7
                factors = [(1 - config.seasonal_weight) + config.seasonal_weight * value / level for value in weekday_means]
            require(all(isclose(a, b, abs_tol=1e-12) for a, b in zip(factors, audit["factores_lunes_a_domingo"])), f"{sku}: factores semanales no coinciden con ventas")
            residuals = [float(daily.get(day, 0)) - mean * factors[day.dayofweek] for day in dates]
            residual_mean = sum((offset + 1) * residual for offset, residual in enumerate(residuals)) / denominator
            variance_denominator = denominator - sum(weight ** 2 for weight in range(1, len(dates) + 1)) / denominator
            sigma = sqrt(sum((offset + 1) * (residual - residual_mean) ** 2 for offset, residual in enumerate(residuals)) / variance_denominator) if variance_denominator > 0 else 0.0
            require(isclose(sigma, audit["sigma_diaria_sin_redondear"], abs_tol=1e-10), f"{sku}: variabilidad no reproducible")
            lead = int(item.lead_time_dias)
            forecast = sum(mean * factors[(cutoff + pd.Timedelta(days=day)).dayofweek] for day in range(lead))
            require(isclose(forecast, row["pronostico_demanda_lead_time"], abs_tol=0.000001), f"{sku}: pronóstico no coincide con operandos publicados")
            safety = NormalDist().inv_cdf(config.service_level) * sigma * sqrt(lead)
            rop = forecast + safety
            require(isclose(safety, row["stock_seguridad"], abs_tol=0.000001), f"{sku}: stock de seguridad no reproducible")
            require(isclose(rop, row["punto_reorden"], abs_tol=0.000001), f"{sku}: ROP no reproducible")
            horizon = lead + config.review_days
            horizon_demand = sum(mean * factors[(cutoff + pd.Timedelta(days=day)).dayofweek] for day in range(horizon))
            target = float(Decimal(str(horizon_demand + NormalDist().inv_cdf(config.service_level) * sigma * sqrt(horizon))).quantize(Decimal(".000001"), rounding=ROUND_HALF_UP))
            require(isclose(target, row["objetivo_cobertura_unidades"], abs_tol=0.000001), f"{sku}: objetivo de compra diferente de demanda + seguridad")
            arrival = pd.to_datetime(lots.fecha_llegada_transito, errors="coerce")
            eligible_transit = int(lots.loc[arrival.between(cutoff, cutoff + pd.Timedelta(days=horizon)), "stock_en_transito"].sum())
            require(eligible_transit == row["stock_en_transito_a_tiempo"], f"{sku}: tránsito elegible no coincide con ETA")
            if usable is not None:
                position = float(Decimal(str(usable + eligible_transit - row["merma_proyectada_fefo"])).quantize(Decimal(".000001"), rounding=ROUND_HALF_UP))
                require(isclose(position, row["posicion_inventario_horizonte"], abs_tol=0.000001), f"{sku}: posición de inventario inconsistente")
                deficit = max(Decimal(0), Decimal(str(target)) - Decimal(str(position)))
                expected_boxes = max(int(item.moq_cajas), ceil(deficit / case_size)) if deficit > 0 and row["estado_sugerencia"] == "CALCULADO" else 0
                require(boxes == expected_boxes, f"{sku}: cajas diferentes de faltante redondeado y MOQ")
        require(row["estado_salud"] == health["estado_salud"], f"{sku}: estados de salud diferentes")
        expiry_units = int(lots.loc[lots.expiry.lt(cutoff + pd.Timedelta(days=120)), "stock_disponible"].sum())
        require(usd(health["inventario_riesgo_vencimiento_usd"]) == usd(Decimal(expiry_units) * Decimal(str(item.costo_unitario_usd))), f"{sku}: vencimiento no coincide con lotes")
        verified.append({
            "sku_id": sku, "nombre_producto": row["nombre_producto"], "proveedor_id": row["proveedor_id"],
            "stock_excel": available, "stock_api": row["stock_disponible"],
            "transito_excel": transit, "transito_api": row["stock_en_transito"],
            "costo_catalogo_usd": float(item.costo_unitario_usd), "costo_api_usd": row["costo_unitario_usd"],
            "cantidad_cajas": boxes, "unidades_por_caja": case_size, "cantidad_unidades": units,
            "total_calculado_usd": float(expected_cost), "total_api_usd": row["costo_total_usd"],
            "demanda_diaria_calculada": mean, "demanda_diaria_api": row["demanda_diaria_ponderada"],
            "pronostico_calculado": forecast, "pronostico_api": row["pronostico_demanda_lead_time"],
            "seguridad_calculada": safety, "seguridad_api": row["stock_seguridad"],
            "rop_calculado": rop, "rop_api": row["punto_reorden"],
            "cobertura_dias": row["cobertura_dias"], "estado_ocs": row["estado_salud"], "estado_dashboard": health["estado_salud"], "resultado": "COINCIDE",
        })
    counts = dict(Counter(row["estado_salud"] for row in rows))
    for state, count in order["salud_catalogo"].items():
        require(count == dashboard["salud_catalogo"][state] == counts.get(state, 0), f"Contador {state} distinto del listado")
    for key in ["capital_sobreinventario_usd", "venta_riesgo_quiebre_usd", "inventario_riesgo_vencimiento_usd"]:
        require(usd(dashboard[key]) == sum((usd(row[key]) for row in detail if row[key] is not None), Decimal(0)), f"KPI {key} no suma sus filas")
    for supplier in order["resumen_proveedores"]:
        subset = [row for row in rows if row["proveedor_id"] == supplier["proveedor_id"]]
        require(usd(supplier["costo_total_usd"]) == sum((usd(row["costo_total_usd"]) for row in subset), Decimal(0)), "Total de proveedor inconsistente")
        require(supplier["cantidad_total_cajas"] == sum(row["cantidad_sugerida_cajas"] for row in subset), "Cajas de proveedor inconsistentes")
    buy_lines = [row for row in rows if row["cantidad_sugerida_cajas"] > 0]
    export_bytes = b""
    if buy_lines:
        purchase = PurchaseExport(
            fecha_corte=order["fecha_corte"], huella_datos=order["huella_datos"],
            parametros={key: order["parametros"][key] for key in ["ventana_dias", "nivel_servicio", "peso_estacional", "dias_revision"]},
            lineas=[dict(proveedor_id=row["proveedor_id"], sku_id=row["sku_id"], cantidad_cajas=row["cantidad_sugerida_cajas"]) for row in buy_lines],
        )
        exported = main.export_contents(contents, source_summary(), [], purchase)
        require(exported.status_code == 200, "La exportación rechazó la OCS")
        export_bytes = exported.body
        book = load_workbook(BytesIO(export_bytes), data_only=True)
        lines = list(book["Orden_Compra"].values)
        exported_rows = [dict(zip(lines[0], line)) for line in lines[1:]]
        expected = {row["sku_id"]: row for row in buy_lines}
        require(len(exported_rows) == len(expected) == len({line["sku_id"] for line in exported_rows}), "XLSX pierde o duplica líneas")
        for line in exported_rows:
            require(line["sku_id"] in expected, "XLSX contiene SKU no seleccionado")
            source = expected[line["sku_id"]]
            for exported_key, source_key in [("cantidad_cajas", "cantidad_sugerida_cajas"), ("cantidad_unidades", "cantidad_sugerida_unidades"), ("costo_unitario_usd", "costo_unitario_usd"), ("total_usd", "costo_total_usd")]:
                require(line[exported_key] == source[source_key], f"XLSX {line['sku_id']}: {exported_key} diferente de la API")
        require(sum((usd(line["total_usd"]) for line in exported_rows), Decimal(0)) == total, "Total XLSX diferente del total OCS")
    require(original_hashes == {name: hashlib.sha256(content).hexdigest() for name, content in contents.items()}, "La auditoría alteró las fuentes")
    report = {
        "estado": "verificado", "fecha_corte": order["fecha_corte"], "huella_datos": order["huella_datos"],
        "fuentes": [{"archivo": name + ".xlsx", "sha256": original_hashes[name], "filas": len(raw[name])} for name in main.SCHEMAS],
        "parametros": order["parametros"], "skus_verificados": len(rows), "lineas_compra": len(buy_lines),
        "cajas_compra": sum(row["cantidad_sugerida_cajas"] for row in rows), "total_compra_usd": str(total),
        "salud_catalogo": order["salud_catalogo"],
        "metricas": {key: dashboard[key] for key in ["capital_sobreinventario_usd", "venta_riesgo_quiebre_usd", "inventario_riesgo_vencimiento_usd"]},
        "advertencias_etl": etl["advertencias"],
    }
    return report, verified, export_bytes


def write_report(directory: Path, report: dict, rows: list[dict], export_bytes: bytes):
    directory.mkdir(parents=True, exist_ok=True)
    (directory / "conciliacion_ocs.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    book = Workbook()
    book.remove(book.active)
    sheets = {
        "Resumen": [{"concepto": key, "valor": value} for key, value in report.items() if isinstance(value, (str, int))],
        "SKU_Verificados": rows,
        "Lineas_Compra": [row for row in rows if row["cantidad_unidades"] > 0],
        "Fuentes": report["fuentes"],
        "Parametros": [{"parametro": key, "valor": value} for key, value in report["parametros"].items()],
        "Metricas": [{"indicador": key, "monto_usd": value} for key, value in report["metricas"].items()],
        "Salud": [{"estado": key, "cantidad": value} for key, value in report["salud_catalogo"].items()],
        "Advertencias": [{"codigo": row["codigo"], "mensaje": row["mensaje"], "cantidad": row.get("cantidad")} for row in report["advertencias_etl"]],
    }
    for title, records in sheets.items():
        sheet = book.create_sheet(title)
        if not records:
            continue
        columns = list(records[0])
        sheet.append(columns)
        for record in records:
            sheet.append([record.get(column) for column in columns])
        for cells in sheet.iter_rows(min_row=2):
            for column, cell in zip(columns, cells):
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                if column.endswith("_usd"):
                    cell.number_format = '"$"#,##0.00'
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(55, max(18, max(len(str(cell.value or "")) for cell in column) + 2))
    book.save(directory / "conciliacion_ocs.xlsx")
    if export_bytes:
        (directory / "orden_compra_verificada.xlsx").write_bytes(export_bytes)


def run():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data-dir", type=Path, default=Path("."))
    parser.add_argument("--output-dir", type=Path, default=Path("reports"))
    parser.add_argument("--ventana-dias", type=int, choices=range(30, 61), default=60)
    args = parser.parse_args()
    try:
        contents = {name: (args.data_dir / f"{name}.xlsx").read_bytes() for name in main.SCHEMAS}
        report, rows, export_bytes = reconcile(contents, ReplenishmentConfig(window_days=args.ventana_dias))
        write_report(args.output_dir, report, rows, export_bytes)
    except (AuditFailure, OSError) as exc:
        parser.exit(1, f"Auditoría rechazada: {exc}\n")
    print(f"VERIFICADO: {report['skus_verificados']} SKU, {report['lineas_compra']} líneas de compra, USD {report['total_compra_usd']}. Informes: {args.output_dir}")


if __name__ == "__main__":
    run()
