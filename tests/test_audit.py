"""La conciliación debe detectar salidas corruptas contra libros independientes."""

import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from fastapi.responses import JSONResponse
from openpyxl import load_workbook

from audit import AuditFailure, reconcile, write_report
import main
from test_engine import engine_tables
from test_ingest import uploads


class AuditTests(unittest.TestCase):
    def setUp(self):
        self.contents = {name: file[1] for name, file in uploads(engine_tables()).items()}

    def test_reconciles_excel_api_and_export_and_writes_reviewable_report(self):
        report, rows, exported = reconcile(self.contents)
        self.assertEqual(report["total_compra_usd"], "48.00")
        self.assertEqual(report["skus_verificados"], 1)
        self.assertEqual(rows[0]["stock_excel"], rows[0]["stock_api"])
        self.assertEqual(rows[0]["resultado"], "COINCIDE")
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            write_report(directory, report, rows, exported)
            workbook = load_workbook(directory / "conciliacion_ocs.xlsx", data_only=True)
            self.assertEqual(workbook["SKU_Verificados"].max_row, 2)
            self.assertEqual(load_workbook(directory / "orden_compra_verificada.xlsx")["Orden_Compra"]["I2"].value, 48)
            self.assertEqual(json.loads((directory / "conciliacion_ocs.json").read_text())["estado"], "verificado")

    def test_rejects_cost_or_stock_corruption_in_api_output(self):
        original = main.replenishment_contents
        for field in ["costo_unitario_usd", "costo_total_usd", "stock_disponible", "stock_en_transito"]:
            with self.subTest(field=field):
                def corrupt(*args, **kwargs):
                    response = original(*args, **kwargs)
                    body = json.loads(response.body)
                    body["orden_sugerida"][0][field] = 99999
                    return JSONResponse(content=body)
                with patch.object(main, "replenishment_contents", side_effect=corrupt):
                    with self.assertRaises(AuditFailure):
                        reconcile(self.contents)

    def test_invalid_sources_do_not_produce_a_successful_reconciliation(self):
        with self.assertRaises(AuditFailure):
            reconcile({**self.contents, "catalogo_sku": b"invalid Excel"})


if __name__ == "__main__":
    unittest.main()
