"""La descarga valida en servidor y exporta cantidades confirmadas, no la IA."""

import copy
import io
import json
import unittest

import httpx
from openpyxl import load_workbook

import main
from test_engine import engine_tables
from test_ingest import uploads


class ExportTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.tables = engine_tables()
        self.files = uploads(self.tables)
        self.client = httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app), base_url="http://test")
        await self.calculate()

    async def asyncTearDown(self):
        await self.client.aclose()

    async def calculate(self):
        response = await self.client.post("/api/v1/replenishment/suggested-order", files=self.files)
        self.assertEqual(response.status_code, 200, response.text)
        self.suggestion = response.json()
        self.order = {
            "fecha_corte": self.suggestion["fecha_corte"],
            "huella_datos": self.suggestion["huella_datos"],
            "parametros": {key: self.suggestion["parametros"][key] for key in ["ventana_dias", "nivel_servicio", "peso_estacional", "dias_revision"]},
            "lineas": [{"proveedor_id": "PROV-1", "sku_id": "SKU-1", "cantidad_cajas": 6}],
        }

    async def export(self, order=None):
        return await self.client.post("/api/v1/replenishment/export", files=self.files, data={"orden": json.dumps(self.order if order is None else order)})

    async def test_confirmed_quantities_real_costs_and_audit(self):
        row = self.suggestion["orden_sugerida"][0]
        for key in ["categoria", "marca", "clasificacion_abc", "costo_unitario_usd", "sku_proveedor"]:
            self.assertIn(key, row)
        response = await self.export()
        self.assertEqual(response.status_code, 200, response.text[:200] if response.status_code != 200 else "")
        self.assertIn("spreadsheetml", response.headers["Content-Type"])
        self.assertIn("OCS_2026-10-21.xlsx", response.headers["Content-Disposition"])
        book = load_workbook(io.BytesIO(response.content), data_only=False)
        self.assertEqual(book.sheetnames, ["Orden_Compra", "Auditoria", "Parametros"])
        line = list(book["Orden_Compra"].values)[1]
        self.assertEqual(line[4], 6)  # original IA = 4; comprador = 6
        self.assertEqual(line[6], 72)  # 6 cajas de 12
        self.assertEqual(line[7], 1)  # costo del catálogo
        self.assertEqual(line[8], 72)
        self.assertEqual(list(book["Auditoria"].values)[1][2:], (4, 6, 2))

    async def test_invalid_quantities_rejected(self):
        for value in [-1, 2.5, True, "6", 1_000_001]:
            with self.subTest(value=value):
                order = copy.deepcopy(self.order)
                order["lineas"][0]["cantidad_cajas"] = value
                response = await self.export(order)
                self.assertEqual(response.status_code, 422)

    async def test_manual_purchase_on_zero_suggestion_uses_catalog_cost(self):
        self.tables["inventario_actual"].loc[0, "stock_disponible"] = 500
        self.files = uploads(self.tables)
        await self.calculate()
        self.assertEqual(self.suggestion["orden_sugerida"][0]["cantidad_sugerida_cajas"], 0)
        response = await self.export()
        self.assertEqual(response.status_code, 200)
        book = load_workbook(io.BytesIO(response.content))
        self.assertEqual(book["Orden_Compra"]["I2"].value, 72)
        self.assertEqual(book["Auditoria"]["C2"].value, 0)

    async def test_moq_and_empty_purchase(self):
        for value, code in [(1, "MOQ_INSUFICIENTE"), (0, "ORDEN_VACIA")]:
            with self.subTest(value=value):
                self.order["lineas"][0]["cantidad_cajas"] = value
                response = await self.export()
                self.assertEqual(response.status_code, 422)
                self.assertEqual(response.json()["errores"][0]["codigo"], code)

    async def test_duplicate_lines(self):
        self.order["lineas"] *= 2
        response = await self.export()
        self.assertEqual(response.status_code, 422)
        self.assertEqual(response.json()["errores"][0]["codigo"], "LINEA_DUPLICADA")

    async def test_wrong_provider_rejected(self):
        self.order["lineas"][0]["proveedor_id"] = "PROV-OTRO"
        response = await self.export()
        self.assertEqual(response.status_code, 422)
        self.assertEqual(response.json()["errores"][0]["codigo"], "SKU_PROVEEDOR_INVALIDO")

    async def test_source_changed_after_calculation(self):
        self.tables["catalogo_sku"].loc[0, "costo_unitario_usd"] = 2
        self.files = uploads(self.tables)
        response = await self.export()
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.json()["errores"][0]["codigo"], "DATOS_CAMBIARON")

    async def test_wrong_cutoff(self):
        self.order["fecha_corte"] = "2026-10-20"
        response = await self.export()
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.json()["errores"][0]["codigo"], "CORTE_CAMBIO")

    async def test_discontinued_cannot_be_exported_by_manual_override(self):
        self.tables["catalogo_sku"].loc[0, "estado"] = "DESCONTINUADO"
        self.files = uploads(self.tables)
        await self.calculate()
        response = await self.export()
        self.assertEqual(response.status_code, 422)
        self.assertEqual(response.json()["errores"][0]["codigo"], "SKU_NO_COMPRABLE")

    async def test_catalog_text_never_becomes_an_excel_formula(self):
        self.tables["catalogo_sku"].loc[0, "nombre_producto"] = '=HYPERLINK("https://example.com","demo")'
        self.files = uploads(self.tables)
        # Pandas escribiría ese texto como fórmula en el fixture; se fuerza string
        # para simular un nombre literal que llega del ERP, no una celda calculada.
        filename, content, mime = self.files["catalogo_sku"]
        source = load_workbook(io.BytesIO(content))
        headers = [cell.value for cell in source["Catalogo"][1]]
        cell = source["Catalogo"].cell(2, headers.index("nombre_producto") + 1)
        cell.data_type = "s"
        buffer = io.BytesIO()
        source.save(buffer)
        self.files["catalogo_sku"] = (filename, buffer.getvalue(), mime)
        await self.calculate()
        response = await self.export()
        self.assertEqual(response.status_code, 200)
        book = load_workbook(io.BytesIO(response.content), data_only=False)
        self.assertEqual(book["Orden_Compra"]["D2"].data_type, "s")


if __name__ == "__main__":
    unittest.main()
